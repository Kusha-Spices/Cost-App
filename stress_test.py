from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import math
from pathlib import Path
import random
import sqlite3
import statistics
import tempfile
import time

from openpyxl import load_workbook

from database_engine import CostingDatabase, DatabaseCostingEngine, as_number, norm_key
import v5_extensions  # noqa: F401
import v6_extensions  # noqa: F401
import v7_extensions  # noqa: F401
import v9_extensions  # noqa: F401


APP_DIR = Path(__file__).parent
WORKBOOK = APP_DIR / "data" / "Cost Sheet 2026 new(3).xlsx"


class StressFailure(AssertionError):
    pass


def check(condition: bool, message: str) -> None:
    if not condition:
        raise StressFailure(message)


def finite_nonnegative(value: float, label: str) -> None:
    check(math.isfinite(float(value)), f"{label} is not finite: {value!r}")
    check(float(value) >= 0, f"{label} is negative: {value!r}")


def build_engine(db_path: Path) -> DatabaseCostingEngine:
    database = CostingDatabase(db_path)
    database.import_workbook(WORKBOOK)
    database.close()
    engine = DatabaseCostingEngine(db_path)
    engine.ensure_v5_schema()
    engine.ensure_v6_schema()
    return engine


def row_inputs(engine: DatabaseCostingEngine) -> list[dict]:
    return [
        dict(row)
        for row in engine.conn.execute(
            """
            SELECT row_no,item,source,shipping_pack,size_g,sale_type,vendor,misc_pct,
                   expected_total_cost,expected_wholesale_cost
            FROM production_rows
            ORDER BY row_no
            """
        ).fetchall()
    ]


def calculate_row(engine: DatabaseCostingEngine, row: dict, cost_mode: str):
    return engine.calculate_pricing(
        item=row["item"],
        source=row["source"],
        shipping_pack=row["shipping_pack"],
        size_g=row["size_g"],
        sale_type=row["sale_type"],
        vendor=row["vendor"],
        misc_pct=None,
        cost_mode=cost_mode,
    )


def test_import_and_integrity(engine: DatabaseCostingEngine) -> dict:
    result = engine.conn.execute("PRAGMA integrity_check").fetchone()[0]
    check(result == "ok", f"SQLite integrity check failed: {result}")
    counts = engine.database_check()["counts"]
    for table in ("master_vlookup", "matrix_values", "production_rows", "new_cost_variations"):
        check(int(counts.get(table) or 0) > 0, f"{table} did not import")
    return counts


def test_imported_recipes(engine: DatabaseCostingEngine) -> dict:
    headers = engine.conn.execute(
        "SELECT product_name,use_for_costing FROM recipe_headers ORDER BY product_name"
    ).fetchall()
    check(len(headers) == 27, f"Expected 27 imported recipes, found {len(headers)}")
    check(all(int(row["use_for_costing"] or 0) == 0 for row in headers), "Imported recipes unexpectedly changed live costing")

    expected = {
        "Premium Biryani Masala": 18,
        "Gun Powder - Without Garlic": 18,
        "Immunity Booster": 18,
        "Black Pepper Powder": 1,
    }
    totals = {}
    for product, expected_rows in expected.items():
        ingredients = engine.recipe_ingredients(product)
        check(len(ingredients) == expected_rows, f"{product} imported {len(ingredients)} rows, expected {expected_rows}")
        total = sum(float(row["qty_g_per_kg"]) for row in ingredients)
        check(abs(total - 1000) < 0.1, f"{product} recipe totals {total:g}g")
        totals[product] = total

    revised = engine.recipe_ingredients("Chilli Oil Spice Blend C Revised")
    revised_total = sum(float(row["qty_g_per_kg"]) for row in revised)
    check(abs(revised_total - 546.6751918158568) < 1e-6, "Incomplete workbook recipe was silently altered")
    revised_header = engine.recipe_header("Chilli Oil Spice Blend C Revised")
    check("WARNING" in str(revised_header.get("notes")), "Incomplete source recipe is not flagged")
    return {"recipes": len(headers), "verified_1000g_samples": totals, "flagged_source_total_g": revised_total}


def test_all_production_rows(engine: DatabaseCostingEngine) -> dict:
    rows = row_inputs(engine)
    successes = 0
    errors: list[dict] = []
    expected_differences: list[dict] = []
    skipped_incomplete: list[dict] = []
    for row in rows:
        for mode, expected_key in (
            ("D2C / Retail full cost with courier shipping", "expected_total_cost"),
            ("Wholesale cost excluding courier shipping", "expected_wholesale_cost"),
        ):
            if row[expected_key] is None:
                skipped_incomplete.append(
                    {"row": row["row_no"], "item": row["item"], "mode": mode}
                )
                continue
            try:
                result = calculate_row(engine, row, mode)
                cost = result.cost
                for field, value in cost.as_dict().items():
                    if isinstance(value, (int, float)):
                        finite_nonnegative(value, f"row {row['row_no']} {field}")
                components_full = (
                    cost.raw_cost
                    + cost.packaging_cost
                    + cost.sticker_cost
                    + cost.labour_cost
                    + cost.storage_cost
                    + cost.shipping_packing_cost
                    + cost.new_shipping_cost
                    + cost.miscellaneous_cost
                )
                check(
                    abs(components_full - cost.total_cost_with_new_shipping) < 1e-8,
                    f"Full-cost components do not add for row {row['row_no']}",
                )
                components_wholesale = (
                    cost.raw_cost
                    + cost.packaging_cost
                    + cost.sticker_cost
                    + cost.labour_cost
                    + cost.storage_cost
                    + cost.shipping_packing_cost
                    + cost.wholesaler_miscellaneous_cost
                )
                check(
                    abs(components_wholesale - cost.wholesaler_cost_excluding_courier_shipping) < 1e-8,
                    f"Wholesale-cost components do not add for row {row['row_no']}",
                )
                expected = row[expected_key]
                difference = float(result.selected_cost) - float(expected)
                if abs(difference) >= 0.01:
                    expected_differences.append(
                        {
                            "row": row["row_no"],
                            "item": row["item"],
                            "mode": mode,
                            "expected": expected,
                            "calculated": result.selected_cost,
                            "difference": difference,
                        }
                    )
                successes += 1
            except Exception as exc:
                errors.append(
                    {
                        "row": row["row_no"],
                        "item": row["item"],
                        "mode": mode,
                        "error": f"{type(exc).__name__}: {exc}",
                    }
                )
    check(not errors, f"{len(errors)} production-row calculations failed; samples={errors[:8]}")
    return {
        "rows": len(rows),
        "calculations": successes,
        "skipped_incomplete": skipped_incomplete,
        "expected_differences": expected_differences,
    }


def test_invalid_inputs(engine: DatabaseCostingEngine) -> dict:
    rejected = 0
    for value in (None, "", "not-a-number", "NaN", "Infinity", float("nan"), float("inf")):
        try:
            as_number(value)
        except (TypeError, ValueError):
            rejected += 1
        else:
            raise StressFailure(f"Invalid numeric input was accepted: {value!r}")

    base = row_inputs(engine)[0]
    for value in (-100, 0, float("nan"), float("inf")):
        try:
            engine.calculate(
                item=base["item"],
                source=base["source"],
                shipping_pack=base["shipping_pack"],
                size_g=value,
                sale_type=base["sale_type"],
                vendor=base["vendor"],
            )
        except (TypeError, ValueError, KeyError):
            rejected += 1
        else:
            raise StressFailure(f"Invalid size was accepted: {value!r}")

    status, _, _, _ = engine.status_from_profit(100, float("nan"))
    check(status != "Good", "NaN cost was classified as Good")
    return {"invalid_values_rejected": rejected}


def test_edit_persistence_and_audit(engine: DatabaseCostingEngine, db_path: Path) -> dict:
    before_audit = len(engine.audit_log_rows(limit=10000))
    original = engine.cell_value("Master Data", "AB32")
    replacement = float(original) + 1.25
    engine.update_cell_value("Master Data", "AB32", replacement)
    check(abs(float(engine.cell_value("Master Data", "AB32")) - replacement) < 1e-9, "Cell edit did not persist")

    first_mrp = engine.conn.execute(
        "SELECT item,item_norm,size_g,mrp FROM new_cost_variations WHERE mrp IS NOT NULL LIMIT 1"
    ).fetchone()
    new_mrp = float(first_mrp["mrp"]) + 3
    engine.update_mrp_variation(first_mrp["item_norm"], first_mrp["size_g"], new_mrp)
    check(
        abs(float(engine.mrp_for(first_mrp["item"], first_mrp["size_g"])) - new_mrp) < 1e-9,
        "MRP edit did not persist",
    )

    engine.save_recipe(
        "Stress Test Blend",
        [{"ingredient_name": "Black Pepper", "qty_g_per_kg": 600}, {"ingredient_name": "Cumin Seeds", "qty_g_per_kg": 400}],
        "Automated stress test",
    )
    recipe = engine.recipe_ingredients("Stress Test Blend")
    check(len(recipe) == 2, "Recipe save did not persist")
    check(abs(sum(float(r["qty_g_per_kg"]) for r in recipe) - 1000) < 1e-9, "Recipe total changed")

    after_audit = len(engine.audit_log_rows(limit=10000))
    check(after_audit >= before_audit + 3, "Expected edits were not written to audit log")

    engine.db.close()
    reopened = DatabaseCostingEngine(db_path)
    reopened.ensure_v5_schema()
    reopened.ensure_v6_schema()
    check(abs(float(reopened.cell_value("Master Data", "AB32")) - replacement) < 1e-9, "Edit was lost after reopen")
    check(len(reopened.recipe_ingredients("Stress Test Blend")) == 2, "Recipe was lost after reopen")
    return {"audit_entries_added": after_audit - before_audit, "engine": reopened}


def test_clean_rebuild(db_path: Path, engine: DatabaseCostingEngine) -> DatabaseCostingEngine:
    engine.create_product_v6(
        product_name="Stress Test Temporary Product",
        category="Raw Material",
        source_or_type="Test Source",
        wholesale_ex_gst=1,
        retail_with_gst=1.2,
        sizes=[100],
        default_packaging_material="",
    )
    check("Stress Test Temporary Product" in engine.product_list(), "Temporary product setup failed")
    engine.db.close()

    database = CostingDatabase(db_path)
    database.import_workbook(WORKBOOK)
    database.close()
    rebuilt = DatabaseCostingEngine(db_path)
    rebuilt.ensure_v5_schema()
    rebuilt.ensure_v6_schema()
    check(
        "Stress Test Temporary Product" not in rebuilt.product_list(),
        "Rebuild from Excel retained stale app-created product data",
    )
    stale_catalog = rebuilt.conn.execute(
        "SELECT 1 FROM item_catalog WHERE item_norm=?",
        (norm_key("Stress Test Temporary Product"),),
    ).fetchone()
    check(stale_catalog is None, "Rebuild from Excel retained stale item-catalog data")
    stale_recipe = rebuilt.conn.execute(
        "SELECT 1 FROM recipe_headers WHERE product_norm=?",
        (norm_key("Stress Test Blend"),),
    ).fetchone()
    check(stale_recipe is None, "Rebuild from Excel retained stale recipe data")
    return rebuilt


def test_export(engine: DatabaseCostingEngine) -> dict:
    output = engine.export_database_to_excel_v7()
    check(len(output) > 10000, "Excel export is unexpectedly small")
    workbook = load_workbook(BytesIO(output), read_only=True, data_only=True)
    required = {"Import Metadata", "Table Counts", "SKU Margin Report", "Validation Issues"}
    missing = sorted(required - set(workbook.sheetnames))
    check(not missing, f"Excel export is missing sheets: {missing}")
    return {"bytes": len(output), "sheets": workbook.sheetnames}


def _load_worker(db_path: str, samples: list[dict], loops: int, seed: int) -> tuple[int, float]:
    random.seed(seed)
    engine = DatabaseCostingEngine(db_path)
    engine.ensure_v5_schema()
    engine.ensure_v6_schema()
    started = time.perf_counter()
    count = 0
    for _ in range(loops):
        row = random.choice(samples)
        mode = random.choice(engine.cost_mode_options())
        result = calculate_row(engine, row, mode)
        finite_nonnegative(result.selected_cost, "concurrent selected cost")
        count += 1
    engine.db.close()
    return count, time.perf_counter() - started


def test_concurrent_load(engine: DatabaseCostingEngine, db_path: Path) -> dict:
    samples = [
        row
        for row in row_inputs(engine)
        if row["expected_total_cost"] is not None and row["expected_wholesale_cost"] is not None
    ]
    workers = 12
    loops_per_worker = 400
    started = time.perf_counter()
    with ThreadPoolExecutor(max_workers=workers) as pool:
        results = list(
            pool.map(
                lambda args: _load_worker(*args),
                [(str(db_path), samples, loops_per_worker, i) for i in range(workers)],
            )
        )
    elapsed = time.perf_counter() - started
    calculations = sum(r[0] for r in results)
    worker_times = [r[1] for r in results]
    check(calculations == workers * loops_per_worker, "Concurrent calculation count mismatch")
    return {
        "workers": workers,
        "calculations": calculations,
        "wall_seconds": elapsed,
        "calculations_per_second": calculations / elapsed,
        "median_worker_seconds": statistics.median(worker_times),
    }


def run() -> dict:
    report: dict = {"status": "PASS", "tests": {}, "failures": []}
    started = time.perf_counter()
    with tempfile.TemporaryDirectory(prefix="kusha_costing_stress_") as temp:
        db_path = Path(temp) / "stress.sqlite"
        engine = build_engine(db_path)
        tests = [
            ("import_and_integrity", lambda: test_import_and_integrity(engine)),
            ("imported_recipes", lambda: test_imported_recipes(engine)),
            ("all_production_rows", lambda: test_all_production_rows(engine)),
            ("invalid_inputs", lambda: test_invalid_inputs(engine)),
            ("edit_persistence_and_audit", lambda: test_edit_persistence_and_audit(engine, db_path)),
        ]
        for name, fn in tests:
            try:
                result = fn()
                if name == "edit_persistence_and_audit":
                    engine = result.pop("engine")
                report["tests"][name] = result
            except Exception as exc:
                report["status"] = "FAIL"
                report["failures"].append({"test": name, "error": f"{type(exc).__name__}: {exc}"})

        try:
            engine = test_clean_rebuild(db_path, engine)
            report["tests"]["clean_rebuild"] = {"passed": True}
        except Exception as exc:
            report["status"] = "FAIL"
            report["failures"].append({"test": "clean_rebuild", "error": f"{type(exc).__name__}: {exc}"})

        for name, fn in [
            ("excel_export", lambda: test_export(engine)),
            ("concurrent_load", lambda: test_concurrent_load(engine, db_path)),
        ]:
            try:
                report["tests"][name] = fn()
            except Exception as exc:
                report["status"] = "FAIL"
                report["failures"].append({"test": name, "error": f"{type(exc).__name__}: {exc}"})

        try:
            engine.db.close()
        except Exception:
            pass

    report["elapsed_seconds"] = time.perf_counter() - started
    return report


if __name__ == "__main__":
    result = run()
    print(json.dumps(result, indent=2, default=str))
    raise SystemExit(0 if result["status"] == "PASS" else 1)
