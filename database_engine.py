from __future__ import annotations

from dataclasses import dataclass, asdict
from io import BytesIO
import math
from pathlib import Path
import re
import sqlite3
from typing import Any, Dict, List, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


# -----------------------------
# Shared helpers
# -----------------------------

def normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\u00a0", " ").strip()


def norm_key(value: Any) -> str:
    return " ".join(normalize(value).lower().split())


def as_number(value: Any) -> float:
    if value is None or value == "":
        raise ValueError("Expected a number, found blank")
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number):
            raise ValueError(f"Expected a finite number, found {value!r}")
        return number
    text = normalize(value).replace(",", "").lower()
    if text.endswith("g"):
        text = text[:-1].strip()
    number = float(text)
    if not math.isfinite(number):
        raise ValueError(f"Expected a finite number, found {value!r}")
    return number


def finite_number(
    value: Any,
    label: str,
    *,
    minimum: Optional[float] = None,
    maximum: Optional[float] = None,
) -> float:
    number = as_number(value)
    if minimum is not None and number < minimum:
        raise ValueError(f"{label} must be at least {minimum:g}")
    if maximum is not None and number > maximum:
        raise ValueError(f"{label} must be at most {maximum:g}")
    return number


def serialize(value: Any) -> tuple[Optional[str], Optional[float], str]:
    """Return (text, number, kind) for SQLite storage."""
    if value is None:
        return None, None, "blank"
    if isinstance(value, bool):
        return str(value), float(int(value)), "number"
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number):
            return None, None, "blank"
        return None, number, "number"
    return str(value), None, "text"


def deserialize(row: sqlite3.Row | None) -> Any:
    if row is None:
        return None
    kind = row["value_kind"]
    if kind == "blank":
        return None
    if kind == "number":
        return row["value_num"]
    return row["value_text"]


def money(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    return round(float(value), 2)


@dataclass
class CostBreakdown:
    item: str
    source: str
    size_g: float
    sale_type: str
    vendor: str
    shipping_pack: str

    landed_rate_per_g: float
    raw_cost: float

    product_packaging_name: str
    packaging_required_qty: float
    packaging_rate: float
    packaging_cost: float

    sticker_cost: float
    labour_cost: float
    storage_cost: float
    shipping_packing_cost: float
    new_shipping_cost: float
    miscellaneous_cost: float
    total_cost_with_new_shipping: float
    wholesaler_miscellaneous_cost: float = 0.0
    wholesaler_cost_excluding_courier_shipping: float = 0.0

    def as_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class PricingResult:
    cost: CostBreakdown
    mrp: Optional[float]
    profit: Optional[float]
    margin_pct: Optional[float]
    status: str
    status_reason: str
    selected_cost: Optional[float] = None
    cost_mode: str = "D2C / Retail full cost with courier shipping"

    def as_dict(self) -> Dict[str, Any]:
        data = self.cost.as_dict()
        data.update(
            {
                "mrp": self.mrp,
                "profit": self.profit,
                "margin_pct": self.margin_pct,
                "status": self.status,
                "status_reason": self.status_reason,
                "selected_cost": self.selected_cost,
                "cost_mode": self.cost_mode,
            }
        )
        return data


# -----------------------------
# Database creation / import
# -----------------------------

class CostingDatabase:
    """
    SQLite database layer.

    Version 3 keeps Excel as the import source, then calculations read from SQLite.
    This is the safe bridge between the current Excel file and a full website later.
    """

    def __init__(self, db_path: Union[str, Path]):
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(self.db_path, check_same_thread=False, timeout=30)
        self.conn.row_factory = sqlite3.Row
        self.init_schema()

    def close(self) -> None:
        self.conn.close()

    def init_schema(self) -> None:
        cur = self.conn.cursor()
        cur.executescript(
            """
            PRAGMA foreign_keys = ON;

            CREATE TABLE IF NOT EXISTS import_meta (
                key TEXT PRIMARY KEY,
                value TEXT
            );

            CREATE TABLE IF NOT EXISTS cells (
                sheet_name TEXT NOT NULL,
                cell_ref TEXT NOT NULL,
                value_text TEXT,
                value_num REAL,
                value_kind TEXT NOT NULL,
                PRIMARY KEY (sheet_name, cell_ref)
            );

            CREATE TABLE IF NOT EXISTS master_vlookup (
                row_no INTEGER NOT NULL,
                row_key TEXT,
                row_key_norm TEXT,
                col_index INTEGER NOT NULL,
                value_text TEXT,
                value_num REAL,
                value_kind TEXT NOT NULL,
                PRIMARY KEY (row_no, col_index)
            );

            CREATE INDEX IF NOT EXISTS idx_master_vlookup_key_col
            ON master_vlookup(row_key_norm, col_index, row_no);

            CREATE TABLE IF NOT EXISTS matrix_values (
                table_name TEXT NOT NULL,
                row_no INTEGER NOT NULL,
                col_no INTEGER NOT NULL,
                row_key TEXT,
                row_key_norm TEXT,
                col_key TEXT,
                col_key_norm TEXT,
                col_key_num REAL,
                value_text TEXT,
                value_num REAL,
                value_kind TEXT NOT NULL,
                PRIMARY KEY (table_name, row_no, col_no)
            );

            CREATE INDEX IF NOT EXISTS idx_matrix_lookup_text
            ON matrix_values(table_name, row_key_norm, col_key_norm);

            CREATE INDEX IF NOT EXISTS idx_matrix_lookup_num
            ON matrix_values(table_name, row_key_norm, col_key_num);

            CREATE TABLE IF NOT EXISTS production_rows (
                row_no INTEGER PRIMARY KEY,
                item TEXT,
                item_norm TEXT,
                source TEXT,
                shipping_pack TEXT,
                size_g REAL,
                sale_type TEXT,
                vendor TEXT,
                misc_pct REAL,
                expected_total_cost REAL,
                expected_wholesale_cost REAL
            );

            CREATE INDEX IF NOT EXISTS idx_production_item
            ON production_rows(item_norm, row_no);

            CREATE TABLE IF NOT EXISTS in_house_rates (
                row_no INTEGER PRIMARY KEY,
                item TEXT,
                item_norm TEXT,
                rate_per_g REAL
            );

            CREATE INDEX IF NOT EXISTS idx_in_house_item
            ON in_house_rates(item_norm, row_no);

            CREATE TABLE IF NOT EXISTS recipe_headers (
                product_norm TEXT PRIMARY KEY,
                product_name TEXT NOT NULL,
                basis_g REAL DEFAULT 1000,
                active INTEGER DEFAULT 1,
                notes TEXT,
                overhead_pct REAL DEFAULT 0,
                use_for_costing INTEGER DEFAULT 0,
                source_recipe_name TEXT
            );

            CREATE TABLE IF NOT EXISTS recipe_ingredients (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                product_norm TEXT NOT NULL,
                ingredient_norm TEXT NOT NULL,
                ingredient_name TEXT NOT NULL,
                qty_g_per_kg REAL NOT NULL,
                FOREIGN KEY(product_norm) REFERENCES recipe_headers(product_norm) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS new_cost_variations (
                item TEXT NOT NULL,
                item_norm TEXT NOT NULL,
                size_g REAL NOT NULL,
                final_cost REAL,
                mrp REAL,
                profit REAL,
                PRIMARY KEY (item_norm, size_g)
            );

            CREATE TABLE IF NOT EXISTS output_defaults (
                field TEXT PRIMARY KEY,
                value_text TEXT,
                value_num REAL,
                value_kind TEXT NOT NULL
            );
            """
        )
        # Lightweight migrations for existing local databases.
        for stmt in [
            "ALTER TABLE production_rows ADD COLUMN misc_pct REAL",
            "ALTER TABLE production_rows ADD COLUMN expected_wholesale_cost REAL",
            "ALTER TABLE recipe_headers ADD COLUMN overhead_pct REAL DEFAULT 0",
            "ALTER TABLE recipe_headers ADD COLUMN use_for_costing INTEGER DEFAULT 0",
            "ALTER TABLE recipe_headers ADD COLUMN source_recipe_name TEXT",
        ]:
            try:
                cur.execute(stmt)
            except sqlite3.OperationalError:
                pass
        self.conn.commit()

    def clear_import_tables(self) -> None:
        cur = self.conn.cursor()
        for table in [
            "recipe_ingredients",
            "recipe_headers",
            "raw_material_sources",
            "packaging_materials",
            "item_catalog",
            "shipping_zone_weights",
            "audit_log",
            "app_settings",
            "cells",
            "master_vlookup",
            "matrix_values",
            "production_rows",
            "in_house_rates",
            "new_cost_variations",
            "output_defaults",
        ]:
            exists = cur.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
                (table,),
            ).fetchone()
            if exists:
                cur.execute(f"DELETE FROM {table}")
        self.conn.commit()

    def set_meta(self, key: str, value: str) -> None:
        self.conn.execute(
            "INSERT OR REPLACE INTO import_meta(key, value) VALUES (?, ?)",
            (key, value),
        )
        self.conn.commit()

    def meta(self) -> Dict[str, str]:
        rows = self.conn.execute("SELECT key, value FROM import_meta ORDER BY key").fetchall()
        return {row["key"]: row["value"] for row in rows}

    def import_workbook(self, workbook_source: Union[str, Path, bytes, BytesIO]) -> None:
        if isinstance(workbook_source, (str, Path)):
            wb = load_workbook(workbook_source, data_only=True)
            source_name = Path(workbook_source).name
        elif isinstance(workbook_source, bytes):
            wb = load_workbook(BytesIO(workbook_source), data_only=True)
            source_name = "uploaded_workbook.xlsx"
        else:
            wb = load_workbook(workbook_source, data_only=True)
            source_name = "uploaded_workbook.xlsx"

        required = ["Master Data", "Production Cost"]
        for sheet_name in required:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Workbook is missing required sheet: {sheet_name}")

        self.clear_import_tables()
        cur = self.conn.cursor()
        master = wb["Master Data"]
        production = wb["Production Cost"]
        new_cost = wb["New Cost"] if "New Cost" in wb.sheetnames else None
        output = wb["Output Sheet"] if "Output Sheet" in wb.sheetnames else None

        # Store key cells used by the costing engine.
        key_cells = [
            "V127",
            "W127",
            "V126",
            "W126",
            "V111",
            "W111",
            "V110",
            "W110",
            "AF40",
            "AF39",
            "AF38",
            "AB32",
            "AB33",
            "AB34",
            "AB35",
            "AB36",
            "AB37",
            "AB38",
        ]
        for cell_ref in key_cells:
            text, num, kind = serialize(master[cell_ref].value)
            cur.execute(
                "INSERT OR REPLACE INTO cells VALUES (?, ?, ?, ?, ?)",
                ("Master Data", cell_ref, text, num, kind),
            )

        # Master Data B:W as exact VLOOKUP-like data.
        start_col = column_index_from_string("B")
        end_col = column_index_from_string("W")
        for row in range(1, master.max_row + 1):
            key = master.cell(row, start_col).value
            if normalize(key) == "":
                continue
            for col in range(start_col, end_col + 1):
                value = master.cell(row, col).value
                text, num, kind = serialize(value)
                cur.execute(
                    """
                    INSERT OR REPLACE INTO master_vlookup
                    (row_no, row_key, row_key_norm, col_index, value_text, value_num, value_kind)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (row, str(key), norm_key(key), col - start_col + 1, text, num, kind),
                )

        # Matrix tables that replace INDEX/MATCH ranges.
        self._import_matrix(cur, master, "maruthi_packaging", "B202", "L255")
        self._import_matrix(cur, master, "swiss_packaging", "B261", "L314")
        self._import_matrix(cur, master, "packing_required", "B323", "K376")

        # In-house rates from H552:J580, rate in J.
        for row in range(552, 581):
            item = master.cell(row, column_index_from_string("H")).value
            if normalize(item) == "":
                continue
            rate = master.cell(row, column_index_from_string("J")).value
            rate_num = float(rate or 0) if isinstance(rate, (int, float)) else 0.0
            cur.execute(
                "INSERT OR REPLACE INTO in_house_rates VALUES (?, ?, ?, ?)",
                (row, str(item), norm_key(item), rate_num),
            )

        recipes_imported, recipe_rows_imported = self._import_recipes(cur, master)

        # Production rows for product list, default sources and legacy defaults.
        # Columns inferred from current working file: B item, D source, E shipping pack,
        # F size, G type, H vendor, AA total cost with new shipping.
        # User-confirmed duplicate cleanup for this workbook:
        # - Green Cardamom row 75 is an old duplicate; use first occurrence row 4.
        # - Dried Mango Powder [Amchur] row 77 is an old duplicate; use row 62.
        skip_production_rows = {75, 77}
        for row in range(3, production.max_row + 1):
            if row in skip_production_rows:
                continue
            item = production.cell(row, 2).value
            if normalize(item) == "":
                continue
            source = production.cell(row, 4).value
            shipping_pack = production.cell(row, 5).value
            size_val = production.cell(row, 6).value
            sale_type = production.cell(row, 7).value
            vendor = production.cell(row, 8).value
            misc_val = production.cell(row, 10).value  # J
            total_val = production.cell(row, 27).value  # AA full D2C/retail-style cost with new shipping
            wholesale_val = production.cell(row, 29).value  # AC wholesaler cost excluding courier shipping
            try:
                size_num = as_number(size_val)
            except Exception:
                size_num = None
            cur.execute(
                """
                INSERT OR REPLACE INTO production_rows
                (row_no, item, item_norm, source, shipping_pack, size_g, sale_type, vendor, misc_pct, expected_total_cost, expected_wholesale_cost)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    row,
                    str(item),
                    norm_key(item),
                    None if source is None else str(source),
                    None if shipping_pack is None else str(shipping_pack),
                    size_num,
                    None if sale_type is None else str(sale_type),
                    None if vendor is None else str(vendor),
                    float(misc_val) if isinstance(misc_val, (int, float)) else None,
                    float(total_val) if isinstance(total_val, (int, float)) else None,
                    float(wholesale_val) if isinstance(wholesale_val, (int, float)) else None,
                ),
            )

        # New Cost long-format variations for MRP/reporting.
        if new_cost is not None:
            blocks = self._new_cost_size_blocks(new_cost)
            for row in range(3, new_cost.max_row + 1):
                item = new_cost.cell(row, 2).value
                if normalize(item) == "":
                    continue
                for block in blocks:
                    cost_val = new_cost.cell(row, block["cost_col"]).value
                    mrp_val = new_cost.cell(row, block["mrp_col"]).value
                    profit_val = new_cost.cell(row, block["profit_col"]).value
                    cur.execute(
                        """
                        INSERT OR REPLACE INTO new_cost_variations
                        (item, item_norm, size_g, final_cost, mrp, profit)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            str(item),
                            norm_key(item),
                            float(block["size_g"]),
                            float(cost_val) if isinstance(cost_val, (int, float)) else None,
                            float(mrp_val) if isinstance(mrp_val, (int, float)) else None,
                            float(profit_val) if isinstance(profit_val, (int, float)) else None,
                        ),
                    )

        # Output Sheet defaults.
        if output is not None:
            fields = {
                "item": "C5",
                "source": "C6",
                "shipping_pack": "C7",
                "size_g": "C8",
                "sale_type": "C9",
                "vendor": "C10",
                # In the current workbook, C11 stores the expected output amount.
                "expected_amount": "C11",
            }
            for field, cell_ref in fields.items():
                text, num, kind = serialize(output[cell_ref].value)
                cur.execute(
                    "INSERT OR REPLACE INTO output_defaults VALUES (?, ?, ?, ?)",
                    (field, text, num, kind),
                )

        self.set_meta("last_import_source", source_name)
        self.set_meta("db_version", "8.3")
        self.set_meta("recipes_imported", str(recipes_imported))
        self.set_meta("recipe_ingredient_rows", str(recipe_rows_imported))
        self.conn.commit()

    def _import_recipes(self, cur: sqlite3.Cursor, master) -> tuple[int, int]:
        """Import the normalized 1kg recipe blocks embedded in Master Data."""
        product_aliases = {
            "rasam masala": "Rasam Masala Powder",
            "sabji masala": "Sabji Masala Powder",
            "sambar masala": "Sambar Masala Powder",
            "gun powder chutney": "Gun Powder - Without Garlic",
            "multi-grain special dal chutney powder": "Multi-grain Special Chutney Powder",
            "peanut chutney powder": "Peanut Chutney Powder - Without Garlic",
            "pulihora masala": "Pulihora Masala Powder",
            "multigrain metkut powder": "Multi-grain Metkut Powder",
        }
        headers: list[tuple[int, int]] = []
        for row in master.iter_rows(min_row=581):
            for cell in row:
                if norm_key(cell.value) == "ingredients":
                    headers.append((cell.row, cell.column))

        recipes_imported = 0
        ingredient_rows_imported = 0
        for header_row, ingredient_col in headers:
            raw_title = normalize(master.cell(header_row - 1, ingredient_col - 1).value)
            if not raw_title:
                continue
            product_name = product_aliases.get(norm_key(raw_title), raw_title)
            ingredients: list[tuple[str, float]] = []
            overhead_pct = 0.0
            found_total = False

            for row_no in range(header_row + 1, min(master.max_row, header_row + 100) + 1):
                ingredient = normalize(master.cell(row_no, ingredient_col).value)
                ingredient_key = norm_key(ingredient)
                if ingredient_key == "total cost for 1000g":
                    found_total = True
                    break

                overhead_match = re.search(
                    r"energy\s*&\s*miscellaneous\s*cost\s*@\s*([0-9.]+)%",
                    ingredient,
                    flags=re.IGNORECASE,
                )
                if overhead_match:
                    overhead_pct = float(overhead_match.group(1))

                serial = master.cell(row_no, ingredient_col - 1).value
                qty = master.cell(row_no, ingredient_col + 2).value
                if (
                    isinstance(serial, (int, float))
                    and not isinstance(serial, bool)
                    and isinstance(qty, (int, float))
                    and not isinstance(qty, bool)
                    and math.isfinite(float(qty))
                    and float(qty) > 0
                    and ingredient
                ):
                    ingredients.append((ingredient, float(qty)))

            if not ingredients or not found_total:
                continue

            total_qty = sum(qty for _, qty in ingredients)
            notes = f"Imported from Master Data recipe block: {raw_title}."
            if abs(total_qty - 1000.0) > 0.1:
                notes += (
                    f" WARNING: source ingredient quantities total {total_qty:,.3f}g, "
                    "not 1,000g; review before enabling for costing."
                )
            product_key = norm_key(product_name)
            cur.execute(
                """
                INSERT OR REPLACE INTO recipe_headers
                (product_norm, product_name, basis_g, active, notes, overhead_pct, use_for_costing, source_recipe_name)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (product_key, product_name, 1000.0, 1, notes, overhead_pct, 0, raw_title),
            )
            cur.execute("DELETE FROM recipe_ingredients WHERE product_norm=?", (product_key,))
            for ingredient_name, qty in ingredients:
                cur.execute(
                    """
                    INSERT INTO recipe_ingredients
                    (product_norm, ingredient_norm, ingredient_name, qty_g_per_kg)
                    VALUES (?, ?, ?, ?)
                    """,
                    (product_key, norm_key(ingredient_name), ingredient_name, qty),
                )
                ingredient_rows_imported += 1
            recipes_imported += 1

        return recipes_imported, ingredient_rows_imported

    def _import_matrix(self, cur: sqlite3.Cursor, sheet, table_name: str, start_cell: str, end_cell: str) -> None:
        start_col_letters = "".join(ch for ch in start_cell if ch.isalpha())
        start_row = int("".join(ch for ch in start_cell if ch.isdigit()))
        end_col_letters = "".join(ch for ch in end_cell if ch.isalpha())
        end_row = int("".join(ch for ch in end_cell if ch.isdigit()))
        start_col = column_index_from_string(start_col_letters)
        end_col = column_index_from_string(end_col_letters)

        for row in range(start_row + 1, end_row + 1):
            row_key = sheet.cell(row, start_col).value
            if normalize(row_key) == "":
                continue
            for col in range(start_col + 1, end_col + 1):
                col_key = sheet.cell(start_row, col).value
                value = sheet.cell(row, col).value
                text, num, kind = serialize(value)
                try:
                    col_key_num = as_number(col_key)
                except Exception:
                    col_key_num = None
                cur.execute(
                    """
                    INSERT OR REPLACE INTO matrix_values
                    (table_name, row_no, col_no, row_key, row_key_norm, col_key, col_key_norm, col_key_num,
                     value_text, value_num, value_kind)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        table_name,
                        row,
                        col,
                        str(row_key),
                        norm_key(row_key),
                        None if col_key is None else str(col_key),
                        norm_key(col_key),
                        col_key_num,
                        text,
                        num,
                        kind,
                    ),
                )

    def _new_cost_size_blocks(self, ws) -> List[Dict[str, Any]]:
        blocks: List[Dict[str, Any]] = []
        for col in range(1, ws.max_column + 1):
            label = ws.cell(1, col).value
            if not normalize(label):
                continue
            try:
                size = int(as_number(label))
            except Exception:
                continue
            if norm_key(ws.cell(2, col).value) == "new cost":
                blocks.append({"size_g": size, "cost_col": col, "mrp_col": col + 1, "profit_col": col + 2})
        return blocks


# -----------------------------
# Database-backed costing engine
# -----------------------------

class DatabaseCostingEngine:
    """Costing engine that reads from SQLite instead of directly from Excel."""

    def __init__(self, db_path: Union[str, Path]):
        self.db = CostingDatabase(db_path)
        self.conn = self.db.conn

    # -----------------------------
    # DB helpers
    # -----------------------------

    def _single(self, query: str, params: tuple = ()) -> sqlite3.Row | None:
        return self.conn.execute(query, params).fetchone()

    def cell_value(self, sheet_name: str, cell_ref: str) -> Any:
        row = self._single(
            "SELECT value_text, value_num, value_kind FROM cells WHERE sheet_name=? AND cell_ref=?",
            (sheet_name, cell_ref),
        )
        return deserialize(row)

    def vlookup_master(self, key: Any, col_index: int) -> Any:
        row = self._single(
            """
            SELECT value_text, value_num, value_kind
            FROM master_vlookup
            WHERE row_key_norm=? AND col_index=?
            ORDER BY row_no ASC
            LIMIT 1
            """,
            (norm_key(key), int(col_index)),
        )
        if row is None:
            raise KeyError(f"Could not find {key!r} in Master Data VLOOKUP table")
        return deserialize(row)

    def matrix_value(self, table_name: str, row_key: Any, col_key: Any) -> Any:
        row_key_norm = norm_key(row_key)
        try:
            col_num = as_number(col_key)
        except Exception:
            col_num = None

        if col_num is not None:
            row = self._single(
                """
                SELECT value_text, value_num, value_kind
                FROM matrix_values
                WHERE table_name=? AND row_key_norm=? AND col_key_num=?
                ORDER BY row_no ASC, col_no ASC
                LIMIT 1
                """,
                (table_name, row_key_norm, col_num),
            )
            if row is not None:
                return deserialize(row)

        row = self._single(
            """
            SELECT value_text, value_num, value_kind
            FROM matrix_values
            WHERE table_name=? AND row_key_norm=? AND col_key_norm=?
            ORDER BY row_no ASC, col_no ASC
            LIMIT 1
            """,
            (table_name, row_key_norm, norm_key(col_key)),
        )
        if row is None:
            raise KeyError(f"Could not find matrix value {table_name}: {row_key!r}, {col_key!r}")
        return deserialize(row)

    # -----------------------------
    # Lists for UI
    # -----------------------------

    def product_list(self) -> List[str]:
        rows = self.conn.execute(
            """
            SELECT item
            FROM production_rows
            WHERE item IS NOT NULL AND item_norm <> ''
            GROUP BY item_norm
            ORDER BY MIN(row_no)
            """
        ).fetchall()
        return [row["item"] for row in rows]

    def size_options(self) -> List[int]:
        sizes = set()
        for row in self.conn.execute("SELECT DISTINCT col_key_num FROM matrix_values WHERE col_key_num IS NOT NULL"):
            sizes.add(int(row["col_key_num"]))
        for row in self.conn.execute("SELECT DISTINCT size_g FROM new_cost_variations WHERE size_g IS NOT NULL"):
            sizes.add(int(row["size_g"]))
        return sorted(sizes)


    def size_options_for_product(self, item: str, vendor: Optional[str] = None) -> List[int]:
        """Return sizes that are actually mapped for the selected product/vendor.

        This prevents invalid combinations like Black Cardamom + 1g, where the global
        size list contains 1g because another SKU uses it, but the selected product has
        no packaging mapping for that size.
        """
        item_norm = norm_key(item)
        rows = self.conn.execute(
            """
            SELECT col_key_num, value_text, value_num, value_kind
            FROM matrix_values
            WHERE table_name='packing_required'
              AND row_key_norm=?
              AND col_key_num IS NOT NULL
            ORDER BY col_key_num
            """,
            (item_norm,),
        ).fetchall()

        candidate_sizes: List[int] = []
        for row in rows:
            if row["value_kind"] == "blank":
                continue
            value = deserialize(row)
            # In this workbook, unsupported sizes are normally blank. If a formula/import
            # creates a zero, treat it as unsupported as well.
            try:
                if value is None or float(value) <= 0:
                    continue
            except Exception:
                if normalize(value) == "":
                    continue
            candidate_sizes.append(int(row["col_key_num"]))

        if vendor:
            table = "maruthi_packaging" if normalize(vendor) == "Maruthi Plastics" else "swiss_packaging"
            filtered: List[int] = []
            for size in candidate_sizes:
                try:
                    pack_name = self.matrix_value(table, item, size)
                    if normalize(pack_name) and norm_key(pack_name) not in {"none", "na", "n/a"}:
                        filtered.append(size)
                except Exception:
                    pass
            return filtered

        if candidate_sizes:
            return candidate_sizes
        return []

    def vendor_options(self) -> List[str]:
        return ["Maruthi Plastics", "Swiss Pac"]

    def sale_type_options(self) -> List[str]:
        return ["Wholesale", "Retail"]

    def shipping_pack_options(self) -> List[str]:
        seen = set()
        opts: List[str] = []

        def add(value: Any) -> None:
            text = normalize(value)
            key = norm_key(text)
            if text and key not in seen:
                opts.append(text)
                seen.add(key)

        for row in self.conn.execute("SELECT shipping_pack FROM production_rows ORDER BY row_no"):
            add(row["shipping_pack"])

        for row in self.conn.execute("SELECT DISTINCT row_key FROM master_vlookup WHERE col_index=1 ORDER BY row_no"):
            val = row["row_key"]
            key = norm_key(val)
            if any(word in key for word in ["box", "ply", "bubble", "pouch"]):
                add(val)
        return opts

    def source_for_product(self, item: str) -> str:
        row = self._single(
            "SELECT source FROM production_rows WHERE item_norm=? AND source IS NOT NULL ORDER BY row_no LIMIT 1",
            (norm_key(item),),
        )
        if row is not None and normalize(row["source"]):
            return str(row["source"])
        value = self.vlookup_master(item, 2)
        return str(value)

    def default_output_inputs(self) -> Dict[str, Any]:
        rows = self.conn.execute("SELECT field, value_text, value_num, value_kind FROM output_defaults").fetchall()
        return {row["field"]: deserialize(row) for row in rows}

    def cost_mode_options(self) -> List[str]:
        return [
            "D2C / Retail full cost with courier shipping",
            "Wholesale cost excluding courier shipping",
        ]

    def _is_wholesale_cost_mode(self, cost_mode: Optional[str]) -> bool:
        return "wholesale" in norm_key(cost_mode or "") and "excluding" in norm_key(cost_mode or "")

    def misc_pct_for_product(self, item: str, size_g: Optional[Union[int, float, str]] = None, sale_type: Optional[str] = None, vendor: Optional[str] = None) -> float:
        params: List[Any] = [norm_key(item)]
        where = "item_norm=? AND misc_pct IS NOT NULL"
        if size_g is not None:
            try:
                where += " AND size_g=?"
                params.append(float(as_number(size_g)))
            except Exception:
                pass
        if sale_type:
            where += " AND LOWER(TRIM(sale_type))=?"
            params.append(norm_key(sale_type))
        if vendor:
            where += " AND vendor=?"
            params.append(vendor)
        row = self._single(f"SELECT misc_pct FROM production_rows WHERE {where} ORDER BY row_no LIMIT 1", tuple(params))
        if row is None:
            row = self._single("SELECT misc_pct FROM production_rows WHERE item_norm=? AND misc_pct IS NOT NULL ORDER BY row_no LIMIT 1", (norm_key(item),))
        if row is None or row["misc_pct"] is None:
            return 0.10
        val = float(row["misc_pct"] or 0)
        return val / 100.0 if val > 1 else val

    def expected_production_cost_for(self, item: str, size_g: Union[int, float, str], cost_mode: str = "D2C / Retail full cost with courier shipping") -> Optional[float]:
        col = "expected_wholesale_cost" if self._is_wholesale_cost_mode(cost_mode) else "expected_total_cost"
        row = self._single(
            f"SELECT {col} AS expected FROM production_rows WHERE item_norm=? AND size_g=? AND {col} IS NOT NULL ORDER BY row_no LIMIT 1",
            (norm_key(item), float(as_number(size_g))),
        )
        if row is None or row["expected"] is None:
            return None
        return float(row["expected"])

    # -----------------------------
    # Pricing / reports
    # -----------------------------

    def mrp_for(self, item: str, size_g: Union[int, float, str]) -> Optional[float]:
        row = self._single(
            "SELECT mrp FROM new_cost_variations WHERE item_norm=? AND size_g=? LIMIT 1",
            (norm_key(item), float(as_number(size_g))),
        )
        if row is None or row["mrp"] is None:
            return None
        return float(row["mrp"])

    def new_cost_sheet_value_for(self, item: str, size_g: Union[int, float, str]) -> Optional[float]:
        row = self._single(
            "SELECT final_cost FROM new_cost_variations WHERE item_norm=? AND size_g=? LIMIT 1",
            (norm_key(item), float(as_number(size_g))),
        )
        if row is None or row["final_cost"] is None:
            return None
        return float(row["final_cost"])

    @staticmethod
    def status_from_profit(mrp: Optional[float], final_cost: float) -> tuple[str, str, Optional[float], Optional[float]]:
        if mrp is None:
            return "MRP Missing", "No MRP found for this product-size variation", None, None
        try:
            mrp = finite_number(mrp, "MRP", minimum=0)
        except (TypeError, ValueError):
            return "Invalid MRP", "MRP must be a finite non-negative number", None, None
        if mrp <= 0:
            return "MRP Missing", "No MRP found for this product-size variation", None, None
        try:
            final_cost = finite_number(final_cost, "Final cost", minimum=0)
        except (TypeError, ValueError):
            return "Invalid Cost", "Final cost must be a finite non-negative number", None, None
        profit = float(mrp) - float(final_cost)
        margin_pct = profit / float(mrp)
        if profit < 0:
            return "Loss", "Profit is below zero", profit, margin_pct
        if margin_pct < 0.20:
            return "Low Margin", "Margin is below 20%", profit, margin_pct
        return "Good", "Margin is 20% or above", profit, margin_pct

    # -----------------------------
    # Costing components
    # -----------------------------

    def _landed_rate_per_g(self, item: str, source: str, sale_type: str) -> float:
        col = 21 if norm_key(sale_type) == "wholesale" else 22
        if normalize(source) == "In House Production":
            row = self._single(
                "SELECT rate_per_g FROM in_house_rates WHERE item_norm=? ORDER BY row_no LIMIT 1",
                (norm_key(item),),
            )
            if row is None:
                raise KeyError(f"In-house rate missing for {item!r}")
            return float(row["rate_per_g"] or 0)
        value = self.vlookup_master(item, col)
        return float(value or 0)

    def _packaging_required_qty(self, item: str, size_g: float) -> float:
        value = self.matrix_value("packing_required", item, size_g)
        return float(value or 0)

    def _product_packaging_name(self, item: str, size_g: float, vendor: str) -> str:
        table = "maruthi_packaging" if normalize(vendor) == "Maruthi Plastics" else "swiss_packaging"
        value = self.matrix_value(table, item, size_g)
        return str(value)

    def _storage_cost(self, size_g: float) -> float:
        if size_g == 1000:
            return float(self.cell_value("Master Data", "AF40") or 0)
        if size_g in (500, 250):
            return float(self.cell_value("Master Data", "AF39") or 0)
        return float(self.cell_value("Master Data", "AF38") or 0)

    def _shipping_packing_cost(self, shipping_pack: str, size_g: float, sale_type: str) -> float:
        col = 21 if norm_key(sale_type) == "wholesale" else 22
        base_cost = float(self.vlookup_master(shipping_pack, col) or 0)
        if size_g == 1000:
            divisor = 1
        elif size_g == 500:
            divisor = 2
        elif size_g == 250:
            divisor = 3
        elif size_g == 100:
            divisor = 4
        elif size_g == 50:
            divisor = 6
        else:
            divisor = 8
        return base_cost / divisor

    def _new_shipping_cost(self, size_g: float) -> float:
        mapping = {1000: "AB32", 500: "AB33", 250: "AB34", 100: "AB35", 50: "AB36", 25: "AB37"}
        addr = mapping.get(int(size_g), "AB38")
        return float(self.cell_value("Master Data", addr) or 0)

    def calculate(
        self,
        item: str,
        size_g: Union[int, float, str],
        sale_type: str,
        vendor: str,
        shipping_pack: str,
        source: Optional[str] = None,
        misc_pct: Optional[float] = None,
    ) -> CostBreakdown:
        size = finite_number(size_g, "Size", minimum=0.000001)
        source = source if source is not None and normalize(source) else self.source_for_product(item)
        if not normalize(item):
            raise ValueError("Item cannot be blank")
        if not normalize(source):
            raise ValueError(f"Source is missing for {item!r}")
        if not normalize(vendor):
            raise ValueError("Packaging vendor cannot be blank")
        if not normalize(shipping_pack):
            raise ValueError("Shipping packing material cannot be blank")
        sale_type = "Wholesale" if norm_key(sale_type) == "wholesale" else "Retail"
        col = 21 if sale_type == "Wholesale" else 22

        rate = finite_number(
            self._landed_rate_per_g(item, source, sale_type),
            "Landed rate per gram",
            minimum=0,
        )
        raw_cost = size * rate

        packaging_required_qty = finite_number(
            self._packaging_required_qty(item, size),
            "Packaging quantity",
            minimum=0,
        )
        product_packaging_name = self._product_packaging_name(item, size, vendor)
        if not normalize(product_packaging_name) or norm_key(product_packaging_name) in {"none", "na", "n/a"}:
            raise ValueError(f"Packaging material is not mapped for {item!r} / {size:g}g / {vendor}")
        packaging_rate = finite_number(
            self.vlookup_master(product_packaging_name, col) or 0,
            "Packaging rate",
            minimum=0,
        )
        packaging_cost = packaging_rate * packaging_required_qty

        sticker_rate = finite_number(
            (self.cell_value("Master Data", "V127") if sale_type == "Wholesale" else self.cell_value("Master Data", "W127")) or 0,
            "Sticker rate",
            minimum=0,
        )
        labour_rate = finite_number(
            (self.cell_value("Master Data", "V111") if sale_type == "Wholesale" else self.cell_value("Master Data", "W111")) or 0,
            "Labour rate",
            minimum=0,
        )
        sticker_cost = sticker_rate * packaging_required_qty
        labour_cost = labour_rate * packaging_required_qty

        storage_cost = self._storage_cost(size)
        shipping_packing_cost = self._shipping_packing_cost(shipping_pack, size, sale_type)
        new_shipping_cost = self._new_shipping_cost(size)
        if misc_pct is None:
            misc_pct = self.misc_pct_for_product(item, size, sale_type, vendor)
        misc_pct = finite_number(misc_pct, "Miscellaneous percentage", minimum=0, maximum=1)
        # New Excel row-3 logic: full D2C/Retail-style miscellaneous includes every cost element,
        # including storage and courier/new shipping.
        miscellaneous_cost = (raw_cost + packaging_cost + sticker_cost + labour_cost + storage_cost + shipping_packing_cost + new_shipping_cost) * misc_pct
        total = raw_cost + packaging_cost + sticker_cost + labour_cost + storage_cost + shipping_packing_cost + new_shipping_cost + miscellaneous_cost
        # Wholesaler cost excludes the courier/new shipping charge, but includes secondary shipping packing material.
        wholesaler_miscellaneous_cost = (raw_cost + packaging_cost + sticker_cost + labour_cost + storage_cost + shipping_packing_cost) * misc_pct
        wholesaler_cost = raw_cost + packaging_cost + sticker_cost + labour_cost + storage_cost + shipping_packing_cost + wholesaler_miscellaneous_cost

        return CostBreakdown(
            item=item,
            source=source,
            size_g=size,
            sale_type=sale_type,
            vendor=vendor,
            shipping_pack=shipping_pack,
            landed_rate_per_g=rate,
            raw_cost=raw_cost,
            product_packaging_name=product_packaging_name,
            packaging_required_qty=packaging_required_qty,
            packaging_rate=packaging_rate,
            packaging_cost=packaging_cost,
            sticker_cost=sticker_cost,
            labour_cost=labour_cost,
            storage_cost=storage_cost,
            shipping_packing_cost=shipping_packing_cost,
            new_shipping_cost=new_shipping_cost,
            miscellaneous_cost=miscellaneous_cost,
            total_cost_with_new_shipping=total,
            wholesaler_miscellaneous_cost=wholesaler_miscellaneous_cost,
            wholesaler_cost_excluding_courier_shipping=wholesaler_cost,
        )

    def calculate_pricing(
        self,
        item: str,
        size_g: Union[int, float, str],
        sale_type: str,
        vendor: str,
        shipping_pack: str,
        source: Optional[str] = None,
        misc_pct: Optional[float] = None,
        cost_mode: str = "D2C / Retail full cost with courier shipping",
    ) -> PricingResult:
        cost = self.calculate(item=item, size_g=size_g, sale_type=sale_type, vendor=vendor, shipping_pack=shipping_pack, source=source, misc_pct=misc_pct)
        mrp = self.mrp_for(item, size_g)
        selected_cost = cost.wholesaler_cost_excluding_courier_shipping if self._is_wholesale_cost_mode(cost_mode) else cost.total_cost_with_new_shipping
        status, reason, profit, margin_pct = self.status_from_profit(mrp, selected_cost)
        return PricingResult(cost=cost, mrp=mrp, profit=profit, margin_pct=margin_pct, status=status, status_reason=reason, selected_cost=selected_cost, cost_mode=cost_mode)

    def costing_defaults_for_product(self, item: str, size_g: Optional[Union[int, float, str]] = None) -> Dict[str, Any]:
        params: List[Any] = [norm_key(item)]
        size_clause = ""
        if size_g is not None:
            size_clause = "CASE WHEN size_g=? THEN 0 ELSE 1 END,"
            params.append(float(as_number(size_g)))
        row = self._single(
            f"""
            SELECT source, shipping_pack, sale_type, vendor
            FROM production_rows
            WHERE item_norm=?
            ORDER BY {size_clause} row_no
            LIMIT 1
            """,
            tuple(params),
        )
        if row is None:
            raise KeyError(f"No costing defaults found for {item!r}")
        return dict(row)

    def sku_report_from_new_cost(self, cost_mode: str = "D2C / Retail full cost with courier shipping") -> List[Dict[str, Any]]:
        """Recalculate portfolio costs from current SQLite master data.

        The old implementation displayed static imported costs. That meant Calculator
        values changed after an edit while Reports stayed stale. This version uses the
        same calculation engine as Calculator and clearly marks incomplete setups.
        """
        rows: List[Dict[str, Any]] = []
        for row in self.conn.execute(
            """
            SELECT item, size_g, mrp
            FROM new_cost_variations
            WHERE mrp IS NOT NULL AND mrp > 0
            ORDER BY item, size_g
            """
        ):
            mrp = float(row["mrp"])
            try:
                defaults = self.costing_defaults_for_product(row["item"], row["size_g"])
                result = self.calculate_pricing(
                    item=row["item"],
                    size_g=row["size_g"],
                    source=defaults.get("source"),
                    shipping_pack=defaults.get("shipping_pack"),
                    sale_type=defaults.get("sale_type") or "Wholesale",
                    vendor=defaults.get("vendor") or "Maruthi Plastics",
                    misc_pct=None,
                    cost_mode=cost_mode,
                )
                rows.append(
                    {
                        "Item": row["item"],
                        "Size (g)": int(row["size_g"]),
                        "Final Cost": round(float(result.selected_cost), 2),
                        "MRP": round(mrp, 2),
                        "Profit": round(float(result.profit), 2) if result.profit is not None else None,
                        "Margin %": result.margin_pct,
                        "Status": result.status,
                        "Status Reason": result.status_reason,
                        "Cost Mode": cost_mode,
                    }
                )
            except Exception as exc:
                rows.append(
                    {
                        "Item": row["item"],
                        "Size (g)": int(row["size_g"]),
                        "Final Cost": None,
                        "MRP": round(mrp, 2),
                        "Profit": None,
                        "Margin %": None,
                        "Status": "Setup Incomplete",
                        "Status Reason": str(exc),
                        "Cost Mode": cost_mode,
                    }
                )
        return rows

    def validate_output_sheet(self) -> Dict[str, Any]:
        defaults = self.default_output_inputs()
        if not defaults:
            return {"ok": False, "message": "No Output Sheet defaults found in database."}

        calc = self.calculate(
            item=defaults.get("item"),
            source=defaults.get("source"),
            shipping_pack=defaults.get("shipping_pack"),
            size_g=defaults.get("size_g"),
            sale_type=defaults.get("sale_type"),
            vendor=defaults.get("vendor"),
        )
        expected = self.expected_production_cost_for(defaults.get("item"), defaults.get("size_g"), "D2C / Retail full cost with courier shipping")
        difference = None
        ok = None
        if isinstance(expected, (int, float)):
            difference = calc.total_cost_with_new_shipping - expected
            ok = abs(difference) < 0.01
        return {
            "ok": ok,
            "expected": expected,
            "calculated": calc.total_cost_with_new_shipping,
            "difference": difference,
            "inputs": defaults,
            "breakdown": calc.as_dict(),
        }


    # -----------------------------
    # Editable master data helpers (Version 4)
    # -----------------------------

    def update_cell_value(self, sheet_name: str, cell_ref: str, value: Any) -> None:
        text, num, kind = serialize(value)
        self.conn.execute(
            """
            INSERT OR REPLACE INTO cells(sheet_name, cell_ref, value_text, value_num, value_kind)
            VALUES (?, ?, ?, ?, ?)
            """,
            (sheet_name, cell_ref, text, num, kind),
        )
        self.conn.commit()

    def master_value_by_row(self, row_no: int, col_index: int) -> Any:
        row = self._single(
            """
            SELECT value_text, value_num, value_kind
            FROM master_vlookup
            WHERE row_no=? AND col_index=?
            """,
            (int(row_no), int(col_index)),
        )
        return deserialize(row)

    def update_master_value_by_row(self, row_no: int, col_index: int, value: Any) -> None:
        existing = self._single(
            "SELECT row_key, row_key_norm FROM master_vlookup WHERE row_no=? LIMIT 1",
            (int(row_no),),
        )
        if existing is None:
            raise KeyError(f"Master Data row {row_no} not found")
        text, num, kind = serialize(value)
        self.conn.execute(
            """
            INSERT OR REPLACE INTO master_vlookup
            (row_no, row_key, row_key_norm, col_index, value_text, value_num, value_kind)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (int(row_no), existing["row_key"], existing["row_key_norm"], int(col_index), text, num, kind),
        )
        self.conn.commit()

    def list_master_rate_rows(self, search: str = "", limit: int = 500) -> List[Dict[str, Any]]:
        params: list[Any] = []
        where = "WHERE k.col_index=1"
        if normalize(search):
            where += " AND k.row_key_norm LIKE ?"
            params.append(f"%{norm_key(search)}%")
        params.append(int(limit))
        rows = self.conn.execute(
            f"""
            SELECT k.row_no, k.row_key,
                   w.value_num AS wholesale_num, w.value_text AS wholesale_text, w.value_kind AS wholesale_kind,
                   r.value_num AS retail_num, r.value_text AS retail_text, r.value_kind AS retail_kind,
                   s.value_text AS source_text, s.value_num AS source_num, s.value_kind AS source_kind
            FROM master_vlookup k
            LEFT JOIN master_vlookup w ON w.row_no=k.row_no AND w.col_index=21
            LEFT JOIN master_vlookup r ON r.row_no=k.row_no AND r.col_index=22
            LEFT JOIN master_vlookup s ON s.row_no=k.row_no AND s.col_index=2
            {where}
            ORDER BY k.row_no
            LIMIT ?
            """,
            tuple(params),
        ).fetchall()
        out: List[Dict[str, Any]] = []
        for row in rows:
            out.append({
                "row_no": row["row_no"],
                "name": row["row_key"],
                "source_or_type": deserialize(row) if False else self.master_value_by_row(row["row_no"], 2),
                "wholesale": deserialize({"value_num": row["wholesale_num"], "value_text": row["wholesale_text"], "value_kind": row["wholesale_kind"]}) if row["wholesale_kind"] is not None else None,
                "retail": deserialize({"value_num": row["retail_num"], "value_text": row["retail_text"], "value_kind": row["retail_kind"]}) if row["retail_kind"] is not None else None,
            })
        return out

    def packaging_rate_rows(self, search: str = "", limit: int = 500) -> List[Dict[str, Any]]:
        # Build a relevant packaging/material list from all matrix values and common shipping pack options.
        names: set[str] = set()
        for row in self.conn.execute(
            """
            SELECT value_text AS name FROM matrix_values
            WHERE table_name IN ('maruthi_packaging', 'swiss_packaging')
              AND value_text IS NOT NULL AND TRIM(value_text) <> ''
            UNION
            SELECT shipping_pack AS name FROM production_rows
            WHERE shipping_pack IS NOT NULL AND TRIM(shipping_pack) <> ''
            """
        ):
            name = normalize(row["name"])
            if name and norm_key(name) not in {"none", "na", "n/a"}:
                names.add(name)
        query = norm_key(search)
        out: List[Dict[str, Any]] = []
        for name in sorted(names, key=lambda x: x.lower()):
            if query and query not in norm_key(name):
                continue
            key = norm_key(name)
            row = self._single(
                "SELECT row_no FROM master_vlookup WHERE row_key_norm=? ORDER BY row_no LIMIT 1",
                (key,),
            )
            if row is None:
                continue
            row_no = int(row["row_no"])
            out.append({
                "row_no": row_no,
                "name": name,
                "wholesale": self.master_value_by_row(row_no, 21),
                "retail": self.master_value_by_row(row_no, 22),
            })
            if len(out) >= limit:
                break
        return out

    def matrix_table_rows(self, table_name: str, item_filter: str = "", limit: int = 500) -> List[Dict[str, Any]]:
        where = "WHERE table_name=?"
        params: list[Any] = [table_name]
        if normalize(item_filter):
            where += " AND row_key_norm LIKE ?"
            params.append(f"%{norm_key(item_filter)}%")
        params.append(int(limit))
        rows = self.conn.execute(
            f"""
            SELECT row_no, col_no, row_key, col_key, col_key_num, value_text, value_num, value_kind
            FROM matrix_values
            {where}
            ORDER BY row_no, col_key_num, col_no
            LIMIT ?
            """,
            tuple(params),
        ).fetchall()
        return [
            {
                "row_no": r["row_no"],
                "col_no": r["col_no"],
                "item": r["row_key"],
                "size_g": int(r["col_key_num"]) if r["col_key_num"] is not None else r["col_key"],
                "value": deserialize(r),
            }
            for r in rows
        ]

    def update_matrix_value(self, table_name: str, row_no: int, col_no: int, value: Any) -> None:
        existing = self._single(
            """
            SELECT row_key, row_key_norm, col_key, col_key_norm, col_key_num
            FROM matrix_values
            WHERE table_name=? AND row_no=? AND col_no=?
            """,
            (table_name, int(row_no), int(col_no)),
        )
        if existing is None:
            raise KeyError(f"Matrix cell not found: {table_name} row {row_no} col {col_no}")
        text, num, kind = serialize(value)
        self.conn.execute(
            """
            INSERT OR REPLACE INTO matrix_values
            (table_name, row_no, col_no, row_key, row_key_norm, col_key, col_key_norm, col_key_num, value_text, value_num, value_kind)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                table_name,
                int(row_no),
                int(col_no),
                existing["row_key"],
                existing["row_key_norm"],
                existing["col_key"],
                existing["col_key_norm"],
                existing["col_key_num"],
                text,
                num,
                kind,
            ),
        )
        self.conn.commit()


    def list_mrp_variations(self, search: str = "", limit: int = 500) -> List[Dict[str, Any]]:
        params: list[Any] = []
        where = "WHERE 1=1"
        if normalize(search):
            where += " AND item_norm LIKE ?"
            params.append(f"%{norm_key(search)}%")
        params.append(int(limit))
        rows = self.conn.execute(
            f"""
            SELECT item, item_norm, size_g, final_cost, mrp, profit
            FROM new_cost_variations
            {where}
            ORDER BY item, size_g
            LIMIT ?
            """,
            tuple(params),
        ).fetchall()
        return [
            {
                "item": r["item"],
                "item_norm": r["item_norm"],
                "size_g": int(r["size_g"]) if r["size_g"] is not None else None,
                "final_cost": r["final_cost"],
                "mrp": r["mrp"],
                "profit": r["profit"],
            }
            for r in rows
        ]

    def update_mrp_variation(self, item_norm: str, size_g: Union[int, float, str], mrp: Any) -> None:
        size = finite_number(size_g, "Size", minimum=0.000001)
        current = self._single(
            "SELECT final_cost FROM new_cost_variations WHERE item_norm=? AND size_g=?",
            (item_norm, size),
        )
        if current is None:
            raise KeyError(f"MRP variation not found for {item_norm} / {size_g}g")
        mrp_num = finite_number(mrp, "MRP", minimum=0) if normalize(mrp) else None
        profit = None
        if mrp_num is not None and current["final_cost"] is not None:
            profit = mrp_num - float(current["final_cost"])
        self.conn.execute(
            """
            UPDATE new_cost_variations
            SET mrp=?, profit=?
            WHERE item_norm=? AND size_g=?
            """,
            (mrp_num, profit, item_norm, size),
        )
        self.conn.commit()

    def update_import_meta(self, key: str, value: str) -> None:
        self.conn.execute(
            "INSERT OR REPLACE INTO import_meta(key, value) VALUES (?, ?)",
            (key, value),
        )
        self.conn.commit()

    def database_check(self) -> Dict[str, Any]:
        counts = {}
        for table in [
            "master_vlookup",
            "matrix_values",
            "production_rows",
            "in_house_rates",
            "new_cost_variations",
            "output_defaults",
        ]:
            counts[table] = self.conn.execute(f"SELECT COUNT(*) AS n FROM {table}").fetchone()["n"]
        return {"meta": self.db.meta(), "counts": counts}
